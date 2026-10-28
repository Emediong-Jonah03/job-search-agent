from html import escape
from io import BytesIO
from pathlib import Path
from string import Template

from app.documents.builder import render_cover_letter_text
from app.documents.models import ApplicationDocumentContent, CVDocument, CoverLetter
from app.documents.claim_validation import validate_document_claims
from app.domain.models import Profile

TEMPLATE_ROOT = Path(__file__).with_name("templates")


def _fact(claim) -> str:
    return claim.text if claim else ""


def cv_plain_text(cv: CVDocument) -> str:
    lines: list[str] = []
    if cv.full_name:
        lines.append(cv.full_name.text)
    if cv.contact:
        lines.append(" | ".join(claim.text for claim in cv.contact))
    if cv.links:
        lines.append(" | ".join(claim.text for claim in cv.links))
    if cv.skills:
        lines.extend(["", "Skills", ", ".join(claim.text for claim in cv.skills)])
    if cv.experience:
        lines.extend(["", "Experience"])
        for item in cv.experience:
            heading = " — ".join(text for text in (_fact(item.title), _fact(item.employer)) if text)
            dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
            if heading:
                lines.append(heading)
            if dates:
                lines.append(dates)
            lines.extend(f"• {claim.text}" for claim in item.responsibilities)
    if cv.projects:
        lines.extend(["", "Projects"])
        for item in cv.projects:
            if item.name:
                lines.append(item.name.text)
            if item.description:
                lines.append(item.description.text)
            if item.technologies:
                lines.append("Technologies: " + ", ".join(claim.text for claim in item.technologies))
            lines.extend(f"• {claim.text}" for claim in item.results)
    if cv.education:
        lines.extend(["", "Education"])
        for item in cv.education:
            line = " — ".join(text for text in (_fact(item.qualification), _fact(item.field), _fact(item.institution)) if text)
            if line:
                lines.append(line)
            dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
            if dates:
                lines.append(dates)
    if cv.certifications:
        lines.extend(["", "Certifications"])
        for item in cv.certifications:
            line = " — ".join(text for text in (_fact(item.name), _fact(item.issuer), _fact(item.date)) if text)
            if line:
                lines.append(line)
    return "\n".join(lines).strip()


def _cv_body(cv: CVDocument) -> str:
    sections: list[str] = []
    header = f"<h1>{escape(_fact(cv.full_name))}</h1>" if cv.full_name else ""
    if cv.contact:
        header += f"<p>{escape(' | '.join(claim.text for claim in cv.contact))}</p>"
    if cv.links:
        header += f"<p>{escape(' | '.join(claim.text for claim in cv.links))}</p>"
    sections.append(header)
    if cv.skills:
        sections.append("<section><h2>Skills</h2><p>" + escape(", ".join(claim.text for claim in cv.skills)) + "</p></section>")
    if cv.experience:
        body = ["<section><h2>Experience</h2>"]
        for item in cv.experience:
            heading = " — ".join(text for text in (_fact(item.title), _fact(item.employer)) if text)
            dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
            if heading:
                body.append(f"<h3>{escape(heading)}</h3>")
            if dates:
                body.append(f"<p>{escape(dates)}</p>")
            if item.responsibilities:
                body.append("<ul>" + "".join(f"<li>{escape(claim.text)}</li>" for claim in item.responsibilities) + "</ul>")
        body.append("</section>")
        sections.append("".join(body))
    if cv.projects:
        body = ["<section><h2>Projects</h2>"]
        for item in cv.projects:
            if item.name:
                body.append(f"<h3>{escape(item.name.text)}</h3>")
            if item.description:
                body.append(f"<p>{escape(item.description.text)}</p>")
            if item.technologies:
                body.append(f"<p>Technologies: {escape(', '.join(claim.text for claim in item.technologies))}</p>")
            if item.results:
                body.append("<ul>" + "".join(f"<li>{escape(claim.text)}</li>" for claim in item.results) + "</ul>")
        body.append("</section>")
        sections.append("".join(body))
    if cv.education:
        body = ["<section><h2>Education</h2>"]
        for item in cv.education:
            line = " — ".join(text for text in (_fact(item.qualification), _fact(item.field), _fact(item.institution)) if text)
            dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
            if line:
                body.append(f"<p>{escape(line)}{(' — ' + escape(dates)) if dates else ''}</p>")
        body.append("</section>")
        sections.append("".join(body))
    if cv.certifications:
        body = ["<section><h2>Certifications</h2>"]
        for item in cv.certifications:
            line = " — ".join(text for text in (_fact(item.name), _fact(item.issuer), _fact(item.date)) if text)
            if line:
                body.append(f"<p>{escape(line)}</p>")
        body.append("</section>")
        sections.append("".join(body))
    return "".join(sections)


def _cover_body(letter: CoverLetter) -> str:
    text = render_cover_letter_text(letter)
    paragraphs = text.split("\n\n")
    body = []
    for paragraph in paragraphs:
        if paragraph.startswith("• ") or "\n• " in paragraph:
            items = paragraph.splitlines()
            body.append("<ul>" + "".join(f"<li>{escape(item[2:])}</li>" for item in items if item.startswith("• ")) + "</ul>")
        elif "\n" in paragraph:
            body.append("<p>" + "<br>".join(escape(line) for line in paragraph.splitlines()) + "</p>")
        else:
            body.append(f"<p>{escape(paragraph)}</p>")
    return "".join(body)


def render_html(cv: CVDocument, letter: CoverLetter) -> tuple[str, str]:
    cv_template = Template((TEMPLATE_ROOT / "cv.html").read_text(encoding="utf-8"))
    letter_template = Template((TEMPLATE_ROOT / "cover_letter.html").read_text(encoding="utf-8"))
    return cv_template.safe_substitute(body=_cv_body(cv)), letter_template.safe_substitute(body=_cover_body(letter))


def _render_pdf(html: str, plain_text: str) -> bytes:
    try:
        from weasyprint import HTML
    except ImportError:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape as xml_escape

        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
        styles = getSampleStyleSheet()
        story = []
        for line in plain_text.splitlines():
            if not line.strip():
                story.append(Spacer(1, 5))
            elif line in {"Skills", "Experience", "Projects", "Education", "Certifications"}:
                story.append(Paragraph(f"<b>{xml_escape(line)}</b>", styles["Heading2"]))
            else:
                story.append(Paragraph(xml_escape(line), styles["BodyText"]))
        doc.build(story)
        return output.getvalue()
    return HTML(string=html).write_pdf()


def _render_docx(content: ApplicationDocumentContent, *, cover: bool) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    if cover:
        for paragraph in render_cover_letter_text(content.cover_letter).split("\n\n"):
            if paragraph.startswith("• ") or "\n• " in paragraph:
                for line in paragraph.splitlines():
                    if line.startswith("• "):
                        document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(paragraph)
    else:
        cv = content.cv
        if cv.full_name:
            document.add_heading(cv.full_name.text, 0)
        if cv.contact:
            document.add_paragraph(" | ".join(claim.text for claim in cv.contact))
        if cv.links:
            document.add_paragraph(" | ".join(claim.text for claim in cv.links))
        if cv.skills:
            document.add_heading("Skills", level=1)
            document.add_paragraph(", ".join(claim.text for claim in cv.skills))
        if cv.experience:
            document.add_heading("Experience", level=1)
            for item in cv.experience:
                heading = " — ".join(text for text in (_fact(item.title), _fact(item.employer)) if text)
                if heading:
                    document.add_heading(heading, level=2)
                dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
                if dates:
                    document.add_paragraph(dates)
                for claim in item.responsibilities:
                    document.add_paragraph(claim.text, style="List Bullet")
        if cv.projects:
            document.add_heading("Projects", level=1)
            for item in cv.projects:
                if item.name:
                    document.add_heading(item.name.text, level=2)
                if item.description:
                    document.add_paragraph(item.description.text)
                if item.technologies:
                    document.add_paragraph("Technologies: " + ", ".join(claim.text for claim in item.technologies))
                for claim in item.results:
                    document.add_paragraph(claim.text, style="List Bullet")
        if cv.education:
            document.add_heading("Education", level=1)
            for item in cv.education:
                line = " — ".join(text for text in (_fact(item.qualification), _fact(item.field), _fact(item.institution)) if text)
                dates = " – ".join(text for text in (_fact(item.start_date), _fact(item.end_date)) if text)
                if line or dates:
                    document.add_paragraph(" — ".join(part for part in (line, dates) if part))
        if cv.certifications:
            document.add_heading("Certifications", level=1)
            for item in cv.certifications:
                line = " — ".join(text for text in (_fact(item.name), _fact(item.issuer), _fact(item.date)) if text)
                if line:
                    document.add_paragraph(line)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)


def extract_docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def render_document_files(profile: Profile, content: ApplicationDocumentContent) -> tuple[dict[str, bytes], dict[str, str], str]:
    cv_html, cover_html = render_html(content.cv, content.cover_letter)
    cv_text = cv_plain_text(content.cv)
    cover_text = render_cover_letter_text(content.cover_letter)
    report = validate_document_claims(profile, content, cv_text)
    if not report.passed:
        raise ValueError("Document claim validation failed: " + "; ".join(report.unsupported_claims))
    if len(cover_text.split()) >= 250:
        raise ValueError("Cover letter must contain fewer than 250 words")

    files = {
        "cv.pdf": _render_pdf(cv_html, cv_text),
        "cv.docx": _render_docx(content, cover=False),
        "cover-letter.pdf": _render_pdf(cover_html, cover_text),
        "cover-letter.docx": _render_docx(content, cover=True),
    }
    extracted = {
        "cv.pdf": extract_pdf_text(files["cv.pdf"]),
        "cv.docx": extract_docx_text(files["cv.docx"]),
        "cover-letter.pdf": extract_pdf_text(files["cover-letter.pdf"]),
        "cover-letter.docx": extract_docx_text(files["cover-letter.docx"]),
    }
    for name, text in extracted.items():
        if not text.strip():
            raise ValueError(f"Generated file has no extractable text: {name}")
    for name in ("cv.pdf", "cv.docx"):
        report = validate_document_claims(profile, content, extracted[name])
        if not report.passed:
            raise ValueError(f"Rendered CV failed claim validation ({name}): " + "; ".join(report.unsupported_claims))
    for name in ("cover-letter.pdf", "cover-letter.docx"):
        report = validate_document_claims(
            profile,
            content,
            cv_text,
            cover_text=extracted[name],
        )
        if not report.passed:
            raise ValueError(f"Rendered cover letter failed claim validation ({name}): " + "; ".join(report.unsupported_claims))
    return files, extracted, cv_text

